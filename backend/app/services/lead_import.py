"""
Importing customers from Excel (.xlsx) or CSV.

Owners often have years of enquiries in a register, an Excel sheet or a
Google Sheet ("File → Download → CSV"). Bringing them in means the AI and the
charts know their history from day one, and old customers aren't asked twice.

How it works:
    1. read the file (Excel or CSV, any common separator / encoding)
    2. guess which column is which from the headings ("Mobile No." → phone,
       "Customer" → name, "Converted?" → status, "Amount ₹" → value...)
    3. check every row; explain problems by ROW NUMBER, like the sheet shows them
    4. PREVIEW first (dry_run): nothing is saved, the owner sees what will happen
    5. then import for real. People already in the system (same phone) are skipped.

Imported enquiries keep their original date, so the "busiest days" and "trend"
charts are right. The AI does NOT re-score old leads automatically (real AI
costs money); there's a button for that (/ai/analyze-all).
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timezone
from typing import Optional
from zoneinfo import ZoneInfo

from pydantic import BaseModel, ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Lead, Message
from app.platforms import PLATFORMS
from app.schemas import LeadCreate

IST = ZoneInfo("Asia/Kolkata")
MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024
FIELDS = ("name", "phone", "email", "source", "service", "message", "status", "value", "date", "area")

# Heading words → our field. Headings are compared in lower case without symbols.
SYNONYMS = {
    "name": ["name", "customer", "customer name", "client", "client name", "full name", "lead name", "contact name",
             "patient", "patient name", "student", "member", "member name", "naam"],
    "phone": ["phone", "mobile", "mobile no", "mobile number", "phone no", "phone number", "contact", "contact no",
              "contact number", "whatsapp", "whatsapp number", "number", "cell"],
    "email": ["email", "email id", "email address", "e mail", "mail"],
    "source": ["source", "lead source", "platform", "channel", "came from", "how did you hear", "medium"],
    "service": ["service", "services", "interested in", "requirement", "product", "enquiry for", "package", "plan",
                "treatment", "course", "category"],
    "message": ["message", "notes", "note", "remarks", "comments", "comment", "enquiry", "query", "details", "description"],
    "status": ["status", "stage", "lead status", "result", "outcome", "converted"],
    "value": ["value", "amount", "deal value", "revenue", "price", "fees", "fee", "paid", "sale value", "bill amount"],
    "date": ["date", "enquiry date", "created", "created at", "created on", "lead date", "date added", "timestamp"],
    "area": ["area", "locality", "location", "city", "address", "sector"],
}
STATUS_WORDS = [  # first match wins
    ("lost", ["lost", "not interested", "no", "junk", "spam", "dead", "cancel", "rejected", "closed lost"]),
    ("won", ["won", "converted", "yes", "paid", "joined", "sale", "sold", "closed won", "customer", "done", "booked & paid"]),
    ("booked", ["booked", "appointment", "visit scheduled", "trial booked", "scheduled"]),
    ("quoted", ["quoted", "quote", "proposal", "negotiation", "price sent"]),
    ("contacted", ["contacted", "called", "follow up", "followup", "in progress", "interested", "callback", "pending",
                   "no answer", "not reachable", "ringing", "busy", "call back"]),
    ("new", ["new", "open", "fresh", ""]),
]
SOURCE_WORDS = {"insta": "instagram", "ig": "instagram", "fb": "facebook", "jd": "justdial", "just dial": "justdial",
                "walkin": "walk_in", "walk in": "walk_in", "reference": "referral", "friend": "referral",
                "word of mouth": "referral", "call": "phone_call", "phone": "phone_call", "web": "website",
                "google maps": "google", "gmb": "google", "wa": "whatsapp", "whats app": "whatsapp"}
DATE_FORMATS = ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y", "%Y-%m-%d", "%d %b %Y", "%d %B %Y",
                "%d-%b-%Y", "%d-%b-%y", "%b %d, %Y", "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S",
                "%Y-%m-%dT%H:%M:%S", "%d-%m-%Y %H:%M")


class ImportError_(Exception):
    """A problem with the whole file (not one row)."""


def _norm(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").lower()).strip()


# ---------------------------------------------------------------------------
# 1. Reading the file
# ---------------------------------------------------------------------------
def read_rows(data: bytes, filename: str = "") -> list[list]:
    if len(data) > MAX_BYTES:
        raise ImportError_("The file is bigger than 5 MB. Split it into smaller files.")
    if data[:4] == b"PK\x03\x04":                                 # .xlsx files are zip files inside
        from openpyxl import load_workbook
        try:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as problem:  # noqa: BLE001 - any broken workbook
            raise ImportError_(f"This Excel file couldn't be opened ({problem.__class__.__name__}). Save it again as .xlsx.")
        sheet = book.worksheets[0]
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        book.close()
        return rows
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise ImportError_("This is an old Excel file (.xls). In Excel: File → Save As → 'Excel Workbook (.xlsx)' or CSV.")
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):           # Excel on Windows often saves CSV as cp1252
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if "\x00" in text[:1000]:
        raise ImportError_("This doesn't look like a CSV or Excel file.")
    try:
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _blank(row: list) -> bool:
    return all(cell is None or str(cell).strip() == "" for cell in row)


def guess_columns(headings: list) -> dict[int, str]:
    """{column number: our field}. Exact heading matches first, then 'contains'."""
    mapping: dict[int, str] = {}
    taken: set[str] = set()
    normal = [_norm(h) for h in headings]
    for exact in (True, False):
        for i, heading in enumerate(normal):
            if i in mapping or not heading:
                continue
            for name, words in SYNONYMS.items():
                if name in taken:
                    continue
                if (heading in words) if exact else any(w in heading.split() or (len(w) > 4 and w in heading) for w in words):
                    mapping[i] = name
                    taken.add(name)
                    break
    return mapping


# ---------------------------------------------------------------------------
# 2. Understanding each cell
# ---------------------------------------------------------------------------
def parse_status(value) -> Optional[str]:
    """→ our status, or None if we don't recognise the word (the row is then saved as 'new' with a note)."""
    text = _norm(value)
    for status, words in STATUS_WORDS:
        if text in words or any(len(w) > 3 and w in text for w in words):
            return status
    return None


def parse_source(value) -> str:
    text = _norm(value)
    if not text:
        return "other"
    key = text.replace(" ", "_")
    if key in PLATFORMS:
        return key
    if text in SOURCE_WORDS:
        return SOURCE_WORDS[text]
    for platform_key, platform in PLATFORMS.items():
        if _norm(platform.name) == text or _norm(platform.name).split(" ")[0] == text:
            return platform_key
    for word, platform_key in SOURCE_WORDS.items():
        if word in text:
            return platform_key
    return "other"


def parse_value(value) -> Optional[int]:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)):
        return int(round(value))
    text = str(value).lower().replace(",", "").replace("₹", "").replace("rs.", "").replace("rs", "").replace("inr", "").strip()
    multiplier = 1
    if text.endswith(("k", "thousand")):
        multiplier, text = 1000, re.sub(r"(k|thousand)$", "", text).strip()
    elif text.endswith(("l", "lakh", "lac", "lakhs")):
        multiplier, text = 100_000, re.sub(r"(lakhs|lakh|lac|l)$", "", text).strip()
    try:
        return int(round(float(text) * multiplier))
    except ValueError:
        raise ValueError(f"amount '{value}' isn't a number")


def parse_date(value) -> Optional[datetime]:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        moment = value
    elif isinstance(value, date):
        moment = datetime.combine(value, time(12, 0))
    else:
        text = str(value).strip()
        for fmt in DATE_FORMATS:
            try:
                moment = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        else:
            raise ValueError(f"date '{value}' isn't clear (use DD/MM/YYYY)")
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=IST)                    # dates in the sheet are India time
    moment = moment.astimezone(timezone.utc)
    if moment > datetime.now(timezone.utc):
        raise ValueError(f"date '{value}' is in the future")
    return moment


# ---------------------------------------------------------------------------
# 3-5. Check every row, preview, import
# ---------------------------------------------------------------------------
class RowProblem(BaseModel):
    row: int                       # the row number as Excel shows it
    problem: str


class PreviewRow(BaseModel):
    row: int
    name: Optional[str]
    phone: Optional[str]
    source: str
    service: Optional[str]
    status: str
    value: Optional[int]
    date: Optional[datetime]


class ImportReport(BaseModel):
    dry_run: bool
    file: str
    rows_in_file: int
    ready: int                     # rows that are fine (dry run) / were imported
    duplicates: int                # already in your leads (same phone) or twice in the file
    problems: list[RowProblem]     # rows that can't be imported, and why
    notes: list[RowProblem]        # rows imported with a small guess (e.g. unknown status → new)
    columns: dict[str, str]        # file heading → what we understood
    not_used: list[str]            # headings we didn't use
    preview: list[PreviewRow]


@dataclass
class _Row:
    number: int
    lead: LeadCreate
    status: str
    value: Optional[int]
    when: Optional[datetime]
    extra: dict = field(default_factory=dict)


def run(db: Session, business_id: int, data: bytes, filename: str, dry_run: bool = True,
        columns: Optional[dict[str, str]] = None) -> ImportReport:
    rows = read_rows(data, filename)
    start = next((i for i, row in enumerate(rows) if not _blank(row)), None)
    if start is None:
        raise ImportError_("The file is empty")
    headings = [str(h).strip() if h is not None else "" for h in rows[start]]
    mapping = guess_columns(headings)
    if columns:                                               # the owner corrected our guesses
        for heading, field_name in columns.items():
            if heading in headings and (field_name in FIELDS or field_name == "skip"):
                index = headings.index(heading)
                mapping = {i: f for i, f in mapping.items() if f != field_name}
                if field_name == "skip":
                    mapping.pop(index, None)
                else:
                    mapping[index] = field_name
    if not {"name", "phone"} & set(mapping.values()):
        raise ImportError_("We couldn't find a Name or Phone column. Make sure the first row has headings like "
                           "'Name' and 'Mobile', or download our template.")
    body = [(start + 2 + n, row) for n, row in enumerate(rows[start + 1:]) if not _blank(row)]
    if len(body) > MAX_ROWS:
        raise ImportError_(f"Up to {MAX_ROWS} rows per file, please (this one has {len(body)}). Split it in parts.")

    known_phones = {p for (p,) in db.execute(select(Lead.phone).where(Lead.business_id == business_id,
                                                                        Lead.phone.is_not(None)))}
    seen: set[str] = set()
    good: list[_Row] = []
    problems: list[RowProblem] = []
    notes: list[RowProblem] = []
    duplicates = 0
    for number, row in body:
        cell = {f: row[i] if i < len(row) else None for i, f in mapping.items()}
        text = {f: (str(v).strip() if v is not None else "") for f, v in cell.items()}
        if isinstance(cell.get("phone"), float):               # Excel turns 9811012345 into 9811012345.0
            text["phone"] = str(int(cell["phone"]))
        try:
            status = parse_status(cell.get("status")) if "status" in cell else "new"
            if status is None:
                notes.append(RowProblem(row=number, problem=f"status '{text.get('status')}' not recognised: saved as 'new'"))
                status = "new"
            value = parse_value(cell.get("value"))
            when = parse_date(cell.get("date"))
            details = {"imported_from": filename[:100]}
            if text.get("email"):
                details["email"] = text["email"][:200]
            lead = LeadCreate(business_id=business_id, source=parse_source(cell.get("source")),
                              name=text.get("name")[:100] if text.get("name") else None,
                              phone=text.get("phone") or None, area=(text.get("area") or "")[:100] or None,
                              service=(text.get("service") or "")[:100] or None, details=details,
                              first_message=(text.get("message") or "")[:4000])
        except ValidationError as error:
            problems.append(RowProblem(row=number, problem="; ".join(
                f"{'.'.join(map(str, e['loc'])) or 'row'}: {e['msg'].removeprefix('Value error, ')}" for e in error.errors())))
            continue
        except ValueError as problem:
            problems.append(RowProblem(row=number, problem=str(problem)))
            continue
        if lead.phone and (lead.phone in known_phones or lead.phone in seen):
            duplicates += 1
            continue
        if lead.phone:
            seen.add(lead.phone)
        good.append(_Row(number, lead, status, value if status == "won" else None, when))

    if not dry_run:
        for item in good:
            when = item.when or datetime.now(timezone.utc)
            lead = Lead(**item.lead.model_dump(exclude={"external_id"}), status=item.status, job_value=item.value,
                        created_at=when, updated_at=when, won_at=when if item.status == "won" else None)
            if item.lead.first_message:
                lead.messages.append(Message(direction="inbound", text=item.lead.first_message, sent_at=when))
            db.add(lead)
        db.commit()

    return ImportReport(
        dry_run=dry_run, file=filename, rows_in_file=len(body), ready=len(good), duplicates=duplicates,
        problems=problems[:200], notes=notes[:200], columns={headings[i]: f for i, f in sorted(mapping.items())},
        not_used=[h for i, h in enumerate(headings) if h and i not in mapping],
        preview=[PreviewRow(row=r.number, name=r.lead.name, phone=r.lead.phone, source=r.lead.source,
                            service=r.lead.service, status=r.status, value=r.value, date=r.when) for r in good[:10]])


def template_csv() -> str:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["Name", "Mobile", "Email", "Source", "Service", "Status", "Amount", "Date", "Area", "Notes"])
    writer.writerow(["Priya Sharma", "98110 12345", "priya@example.com", "Instagram", "Zumba", "Won", "₹4,500",
                     "05/10/2026", "Rajouri Garden", "Joined the evening batch"])
    writer.writerow(["Rahul Verma", "9876543210", "", "Justdial", "Personal training", "Contacted", "",
                     "12/10/2026", "Janakpuri", "Asked for weekend timings"])
    return out.getvalue()
